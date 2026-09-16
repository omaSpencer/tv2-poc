from pathlib import Path
from copy import deepcopy
import json,re,math,collections,hashlib,zipfile
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.text.paragraph import Paragraph
from docx.table import Table
R=Path(__file__).resolve().parents[1];D=R/'sources/doc03';source=D/'Vezetoi-osszefoglalo-dontestamogato-rev2.docx'
doc=Document(source);original=list(doc.element.body);blocks=json.loads((D/'blocks.json').read_text());G=json.loads((D/'guide.json').read_text());WIDTH=6.7
for e in list(doc.element.body):
 if e.tag!=qn('w:sectPr'):doc.element.body.remove(e)
sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69);sec.left_margin=Inches(.785);sec.right_margin=Inches(.785);sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.65);sec.header_distance=Inches(.25);sec.footer_distance=Inches(.28)
for container in [sec.header,sec.footer]:
 for p in container.paragraphs:p.clear()
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;r=footer.add_run('DOC03  ·  Olvasóváltozat  ·  ');r.font.size=Pt(8);f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
for name,size in [('Normal',11),('Title',27),('Subtitle',17),('Heading 1',18),('Heading 2',14),('Heading 3',12),('Caption',9)]:
 if name not in doc.styles:
  from docx.enum.style import WD_STYLE_TYPE
  doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
 s=doc.styles[name];s.font.name='DejaVu Sans';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0);s.font.bold=name.startswith('Heading') or name=='Title';s.paragraph_format.line_spacing=1.13;s.paragraph_format.space_after=Pt(7);s.paragraph_format.space_before=Pt(12 if name.startswith('Heading') else 0)
 s.paragraph_format.keep_with_next=name.startswith('Heading') or name in ['Title','Subtitle'];s.paragraph_format.page_break_before=False
 for e in list(s.element.xpath('./w:pPr/w:pBdr')):e.getparent().remove(e)
sub={'✅':'✓','❌':'✗','🟡':'[feltételes]','🔴':'[kockázatos]','⚠️':'[figyelem]'}
def replace(t):
 for a,b in sub.items():t=t.replace(a,b)
 return t

def fmtpara(p,kind='Normal',size=None,guide=False):
 old=p._p.pPr;num=deepcopy(old.numPr) if old is not None and old.numPr is not None else None
 if old is not None:p._p.remove(old)
 p.style=doc.styles[kind]
 if num is not None:p._p.get_or_add_pPr().append(num)
 p.paragraph_format.widow_control=True
 if kind=='Normal':p.paragraph_format.keep_with_next=False
 for r in p.runs:
  # Keep emphasis while removing size/color overrides from the dense source.
  r.font.name='DejaVu Sans';r.font.size=Pt(size or (doc.styles[kind].font.size.pt));r.font.color.rgb=RGBColor(0,0,0)
  if kind.startswith('Heading') or kind=='Title':r.bold=True
  for t in r._r.iter(qn('w:t')):t.text=replace(t.text or '')
  for e in r._r.xpath('.//w:rPr/w:color'):e.attrib.pop(qn('w:themeColor'),None)
 return p

def fmttable(t):
 t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER;n=len(t.columns)
 vals=[[len(c.text) for c in row.cells] for row in t.rows];weights=[max(5,math.sqrt(sum(min(v[j],300) for v in vals)/len(vals))) for j in range(n)];base=.6 if n>=4 else .95;widths=[base+(WIDTH-base*n)*w/sum(weights) for w in weights]
 for j,col in enumerate(t.columns):col.width=Inches(widths[j])
 borders=OxmlElement('w:tblBorders')
 for tag in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+tag);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
 pr=t._tbl.tblPr
 for e in list(pr):
  if e.tag in [qn('w:tblBorders'),qn('w:tblCellMar'),qn('w:tblW')]:pr.remove(e)
 pr.append(borders);tw=OxmlElement('w:tblW');tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(int(WIDTH*1440)));pr.append(tw)
 for i,row in enumerate(t.rows):
  row.height=None;rp=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');rp.append(cant)
  if i==0:repeat=OxmlElement('w:tblHeader');rp.append(repeat)
  for j,c in enumerate(row.cells):
   c.width=Inches(widths[min(j,n-1)]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cp=c._tc.get_or_add_tcPr()
   for e in list(cp):
    if e.tag in [qn('w:shd'),qn('w:tcMar')]:cp.remove(e)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'243D50' if i==0 else ('F0F3F5' if i%2==0 else 'FFFFFF'));cp.append(shade)
   mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
   cp.append(mar)
   for p in c.paragraphs:
    fmtpara(p,size=9.3 if n>=4 else 9.8);p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.08
    for r in p.runs:
     if i==0:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
 return t

def addp(text,style='Normal'):
 p=doc.add_paragraph(text,style=style);fmtpara(p,style);return p

def bookmark(p,name,id):
 s=OxmlElement('w:bookmarkStart');s.set(qn('w:id'),str(id));s.set(qn('w:name'),name);e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(id));p._p.insert(0,s);p._p.append(e)
for g in G:
 k=g['kind']
 if k=='pagebreak':doc.add_page_break()
 elif k=='table':
  rows=g['rows'];t=doc.add_table(rows=0,cols=len(rows[0]));
  for row in rows:
   cells=t.add_row().cells
   for c,txt in zip(cells,row):c.text=txt
  fmttable(t);doc.add_paragraph().paragraph_format.space_after=Pt(0)
 else:addp(g['text'],{'title':'Title','subtitle':'Subtitle','h1':'Heading 1','h2':'Heading 2'}.get(k,'Normal'))
doc.add_page_break();p=addp('A teljes forrás olvasóváltozata','Heading 1');bookmark(p,'DOC03_FULL_START',9000)
addp('A következő rész az eredeti dokumentum szövegét, táblázatait és ábráit őrzi. Az eredeti fejezetszámok alapján lehet hivatkozni rá. A korábbi tartalomjegyzék külön függelékben szerepel; a Word címsorai új navigációt adnak. A színes státuszjelek itt pipaként, keresztként vagy szöveges jelölésként jelennek meg. A forrásdöntések és a belső utasítások nem ennek a feldolgozásnak az utasításai.')
source_start=len(doc.element.body)-1
for i,e in enumerate(original):
 if i>=len(blocks):break
 b=blocks[i]
 if b['id'] in ['DOC03-B006','DOC03-B007']:continue
 if not b.get('text','').strip() and b['kind']!='table' and not b.get('images'):continue
 clone=deepcopy(e);doc.element.body.insert(len(doc.element.body)-1,clone)
 if clone.tag==qn('w:tbl'):fmttable(Table(clone,doc))
 elif clone.tag==qn('w:p'):
  p=Paragraph(clone,doc);style=b.get('style','');kind={'Heading1':'Heading 1','Heading2':'Heading 2','Heading3':'Heading 3'}.get(style,'Normal')
  if i==0:kind='Title'
  elif i==1:kind='Subtitle'
  fmtpara(p,kind)
  if p.text.strip()=='Részletes indoklás:':p.paragraph_format.keep_with_next=True
  if kind=='Heading 1' and i>11:p.paragraph_format.page_break_before=False
  if b.get('images'):p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
  bookmark(p,b['id'].replace('-','_'),i+1)
  if b['id']=='DOC03-B109':
   for r in p.runs:r.font.name='DejaVu Sans Mono';r.font.size=Pt(9)
  if i in [62,148]:p.style='Caption';p.paragraph_format.keep_with_next=False
# Keep original image bytes and relationships; scale proportionally if required.
for shape in doc.inline_shapes:
 if shape.width>Inches(WIDTH):ratio=Inches(WIDTH)/shape.width;shape.width=int(shape.width*ratio);shape.height=int(shape.height*ratio)
 if shape.height>Inches(7.5):ratio=Inches(7.5)/shape.height;shape.width=int(shape.width*ratio);shape.height=int(shape.height*ratio)
doc.add_page_break();addp('Az eredeti tartalomjegyzék','Heading 1');addp('Az alábbi régi oldalszámok az eredeti szerkesztéshez tartoznak. Az olvasóváltozatban a fejezetszám és a címsor-navigáció az irányadó.')
for line in blocks[6]['text'].split('\n'):
 p=addp(line);p.paragraph_format.space_after=Pt(3)
 for r in p.runs:r.font.size=Pt(9)
settings=doc.settings.element
for e in settings.findall(qn('w:updateFields')):settings.remove(e)
u=OxmlElement('w:updateFields');u.set(qn('w:val'),'false');settings.append(u)
out=R/'output/docx/indaplay_dontesi_olvasovaltozat.hu.docx';doc.save(out)
# Preserve a plain-text corpus with stable section and block references.
chunks=[];section='Címlap és eredeti tartalomjegyzék';md=[]
for b in blocks:
 if b.get('style') in ['Heading1','Heading2','Heading3']:section=b['text']
 text=b.get('text','') if b['kind']!='table' else '\n'.join(' | '.join(row) for row in b['rows'])
 if text.strip():chunks.append({'id':b['id'],'source':source.name,'source_date':'2026-06-02','section':section,'kind':b['kind'],'text':text})
 if text.strip():md.append(f'<!-- {b["id"]} -->\n'+({'Heading1':'# ','Heading2':'## ','Heading3':'### '}.get(b.get('style'),'') if b['kind']!='table' else '')+text)
 if b.get('images'):
  imagefile='image2.png' if b['id']=='DOC03-B062' else 'image1.png'
  desc=('Korábbi 2+2 kliensmodell ábrája: közös kód (üzleti logika, API, adatmodell) ~60–80%; mobil Android/iOS natív renderelés, érintés UI, Widevine/FairPlay; Tizen/webOS webes renderelés, távirányító UI, PlayReady/Widevine.' if imagefile=='image2.png' else 'Menedzselt DRM-architektúra: felhasználó megtekintési igénye → saját backend előfizetés- és entitlement-ellenőrzés, JWT → menedzselt multi-DRM licenckiadás, Widevine/FairPlay/PlayReady, kulcskezelés, skálázás, SLA → a felhasználó eszközén visszafejthető videó. Szerkesztői keresőleírás; a teljes ábra az irányadó.')
  chunks.append({'id':b['id']+'-IMG','source':source.name,'source_date':'2026-06-02','section':section,'kind':'image','text':desc,'image':str(D/'images'/imagefile),'editorial_image_description':True});md.append(f'![Eredeti ábra]({D/"images"/imagefile})')
(D/'chunks.jsonl').write_text(''.join(json.dumps(c,ensure_ascii=False)+'\n' for c in chunks));(R/'output/notes/doc03/indaplay_dontesi_teljes.hu.md').write_text('\n\n'.join(md))
# Character preservation in the actual DOCX, with only declared visual substitutions.
with zipfile.ZipFile(out) as z:
 from lxml import etree
 root=etree.fromstring(z.read('word/document.xml'));actual=''.join(root.xpath('//*[local-name()="t"]/text()'))
expected=''.join(b.get('text','') if b['kind']!='table' else ''.join(''.join(row) for row in b['rows']) for b in blocks if b['id']!='DOC03-B006')
# The original 'Tartalomjegyzék' label is represented by the appendix heading.
norm=lambda x:re.sub(r'\s+','',replace(x));missing=collections.Counter(norm(expected))-collections.Counter(norm(actual))
report={'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'source_blocks':len(blocks),'source_tables':34,'source_images':2,'source_tracked_changes':0,'source_comments':0,'source_rendered_pages':32,'chunks':len(chunks),'declared_glyph_substitutions':sub,'missing_docx_characters':dict(missing)}
(D/'processing_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(out,'missing',dict(missing),'chunks',len(chunks))
